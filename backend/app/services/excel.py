"""Build the analytics workbook. Pure function of already-fetched data; runs in a thread."""

import io
from datetime import UTC, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from app.core.time import utcnow
from app.schemas.analytics import CategorySummary, Source, Summary, TrendPoint
from app.schemas.record import CONTROL_CHARACTERS

BOLD = Font(bold=True)


def _styled(sheet, value: object, font: Font) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, value=value)
    cell.font = font
    return cell


def _header(sheet, columns: list[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.append([_styled(sheet, c, BOLD) for c in columns])




def _text(sheet, value: object) -> object:
    """Neutralise formula injection: openpyxl stores text starting with "=" as a formula.

    Such text is written as a string cell instead, so Excel shows it unchanged and never
    evaluates it. Other leading characters (+, -, @) are already stored as plain text in
    .xlsx, so they are left alone; "-40C cold start" stays exactly that. Control characters,
    which .xlsx cannot hold, are removed (the API refuses them; this guards rows stored before).
    """
    if isinstance(value, str):
        value = CONTROL_CHARACTERS.sub("", value)
        if value.startswith("="):
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"  # openpyxl took it for a formula
            return cell
    return value


def _naive(value: datetime) -> datetime:
    return value.replace(tzinfo=None)


def _ms(moment: datetime) -> datetime:
    """Excel keeps milliseconds: cut, not round, so 59.9999 s does not become the next minute."""
    return moment.replace(microsecond=moment.microsecond // 1000 * 1000)


def build_workbook(
    *,
    source: Source,
    all_categories: bool,
    filters: dict[str, str],
    summary: Summary,
    categories: list[CategorySummary],
    trend: list[TrendPoint],
    zone: ZoneInfo,
    rows: list,  # Record or RealtimeReading, matching `source`
    row_limit: int,
    trend_limit: int,
) -> bytes:
    # Write-only mode streams rows instead of keeping a cell object for each: about twice as
    # fast for a 17,000-row Data sheet. Cells must be written in order, styled as created.
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Summary")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 24
    sheet.append([_styled(sheet, "Burn-in monitor analytics report", Font(bold=True, size=14))])
    sheet.append(["Generated (UTC)", utcnow().replace(microsecond=0)])
    sheet.append(["Source", "Test records" if source == Source.records else "Live telemetry"])
    for key, value in filters.items():
        sheet.append([key, _text(sheet, value)])
    sheet.append([])
    sheet.append(["Metric", "Value"])
    sheet.append(["count", summary.count])
    if all_categories:
        # A sum or an average across °C, RPM and W means nothing.
        sheet.append(["total, average, minimum, maximum", "per category on the next sheet: the units differ"])
    else:
        for metric in ("total", "average", "minimum", "maximum"):
            sheet.append([metric, getattr(summary, metric)])
    sheet.append([])
    # Said here, where the reader starts: the statistics cover the whole range, the Data sheet may not.
    if len(rows) >= row_limit:
        sheet.append(["Data sheet", f"Only the latest {row_limit:,} rows; the statistics above cover the whole range"])
    elif summary.count > len(rows):
        sheet.append(["Data sheet", f"{len(rows):,} rows: retention has deleted the other raw readings of this range; "
                                    "the statistics above are kept"])
    else:
        sheet.append(["Data sheet", f"All {len(rows):,} rows"])

    by_cat = workbook.create_sheet("By category")
    _header(by_cat, ["category", "count", "total", "average", "minimum", "maximum"])
    for c in categories:
        by_cat.append([c.category, c.count, c.total, c.average, c.minimum, c.maximum])

    trend_sheet = workbook.create_sheet("Trend")
    _header(trend_sheet, [f"bucket ({zone.key})", "category", "count", "average", "minimum", "maximum"])
    for p in trend:
        local = p.bucket.replace(tzinfo=UTC).astimezone(zone)
        trend_sheet.append([_naive(local), p.category, p.count, p.average, p.minimum, p.maximum])
    if len(trend) >= trend_limit:
        trend_sheet.append([f"Only the newest {trend_limit:,} buckets; choose a longer interval for the whole range"])

    data = workbook.create_sheet("Data")
    if source == Source.records:
        _header(data, ["id", "title", "value", "category", "server", "recorded_at (UTC, to the ms)", "owner"])
        for r in rows:
            data.append([r.id, _text(data, r.title), r.value, r.category, r.server, _ms(r.recorded_at), _text(data, r.owner_username)])
    else:
        _header(data, ["id", "server", "category", "sensor", "value", "threshold", "is_anomaly", "generated_at (UTC, to the ms)"])
        for r in rows:
            data.append([r.id, r.server, r.category, r.sensor, r.value, r.threshold, r.is_anomaly, _ms(r.generated_at)])
    if len(rows) >= row_limit:
        data.append([f"Truncated to the latest {row_limit} rows"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
