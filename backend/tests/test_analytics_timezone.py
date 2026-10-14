"""Independent UTC/local calendar fixtures for historical timezone grouping."""

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from tests.conftest import create_record


@pytest.mark.parametrize("zone,interval,instants,expected", [
    ("Europe/London", "day", ["2026-01-15T23:30:00Z", "2026-01-16T00:30:00Z"],
     ["2026-01-15T00:00:00Z", "2026-01-16T00:00:00Z"]),
    ("America/Ciudad_Juarez", "day", ["2026-01-16T06:30:00Z", "2026-01-16T07:30:00Z"],
     ["2026-01-15T07:00:00Z", "2026-01-16T07:00:00Z"]),
    ("America/Ciudad_Juarez", "day", ["2026-07-16T05:30:00Z", "2026-07-16T06:30:00Z"],
     ["2026-07-15T06:00:00Z", "2026-07-16T06:00:00Z"]),
    ("Europe/London", "hour", ["2026-03-29T00:30:00Z", "2026-03-29T01:30:00Z"],
     ["2026-03-29T00:00:00Z", "2026-03-29T01:00:00Z"]),
    ("Europe/London", "minute", ["2026-10-25T00:30:10Z", "2026-10-25T01:30:10Z"],
     ["2026-10-25T00:30:00Z", "2026-10-25T01:30:00Z"]),
    ("Europe/London", "hour", ["2026-10-25T00:30:00Z", "2026-10-25T01:30:00Z"],
     ["2026-10-25T00:00:00Z", "2026-10-25T01:00:00Z"]),
    ("Australia/Lord_Howe", "hour", ["2026-10-03T15:15:00Z", "2026-10-03T15:45:00Z"],
     ["2026-10-03T14:30:00Z", "2026-10-03T15:30:00Z"]),
    ("Australia/Lord_Howe", "hour", ["2026-04-04T14:45:00Z", "2026-04-04T15:15:00Z"],
     ["2026-04-04T14:00:00Z", "2026-04-04T15:00:00Z"]),
    ("Asia/Kathmandu", "hour", ["2026-01-16T00:10:00Z"], ["2026-01-15T23:15:00Z"]),
    ("Pacific/Chatham", "hour", ["2026-01-16T00:10:00Z"], ["2026-01-15T23:15:00Z"]),
    ("America/Bogota", "day", ["2026-01-16T03:00:00Z"], ["2026-01-15T05:00:00Z"]),
    ("UTC", "minute", ["2026-01-16T00:10:45Z"], ["2026-01-16T00:10:00Z"]),
    ("Pacific/Apia", "day", ["2011-12-29T09:30:00Z", "2011-12-30T10:30:00Z"],
     ["2011-12-28T10:00:00Z", "2011-12-30T10:00:00Z"]),
    ("Africa/Monrovia", "minute", ["1972-01-06T23:59:40Z"], ["1972-01-06T23:59:30Z"]),
    ("America/Sao_Paulo", "day", ["2018-11-04T03:30:00Z"], ["2018-11-04T03:00:00Z"]),
])
async def test_historical_bucket_membership(client, accounts, zone, interval, instants, expected):
    for index, instant in enumerate(instants):
        await create_record(client, accounts.engineer, value=10 * (index + 1), recorded_at=instant)
    response = await client.get("/api/v1/analytics/trend", params={"tz": zone, "interval": interval},
                                headers=accounts.viewer.headers)
    assert response.status_code == 200, response.text
    assert [(p["bucket"], p["count"], p["average"], p["minimum"], p["maximum"]) for p in response.json()] == [
        (bucket, 1, 10 * (index + 1), 10 * (index + 1), 10 * (index + 1))
        for index, bucket in enumerate(expected)
    ]


async def test_transition_day_is_one_bucket_at_historical_midnight(client, accounts):
    # London midnight was GMT; after 01:00 UTC the same local date uses BST.
    for value, instant in [(10, "2026-03-29T00:30:00Z"), (30, "2026-03-29T12:30:00Z")]:
        await create_record(client, accounts.engineer, value=value, recorded_at=instant)
    response = await client.get("/api/v1/analytics/trend", params={"tz": "Europe/London", "interval": "day"},
                                headers=accounts.viewer.headers)
    assert response.json() == [{"bucket": "2026-03-29T00:00:00Z", "category": "cpu_temp",
                                "count": 2, "average": 20.0, "minimum": 10.0, "maximum": 30.0}]


async def test_excel_historical_days_and_alloy_summary(client, accounts):
    for value, instant in [(10, "2026-01-16T06:30:00Z"), (30, "2026-01-16T07:30:00Z")]:
        await create_record(client, accounts.engineer, category="alloy", value=value, recorded_at=instant)
    response = await client.get("/api/v1/analytics/export", headers=accounts.viewer.headers,
                                params={"tz": "America/Ciudad_Juarez", "interval": "day", "category": "alloy"})
    workbook = load_workbook(io.BytesIO(response.content))
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert {key: summary.get(key) for key in ("count", "total", "average", "minimum", "maximum")} == {
        "count": 2, "total": 40, "average": 20, "minimum": 10, "maximum": 30,
    }
    assert list(workbook["Trend"].iter_rows(values_only=True))[1:] == [
        (datetime(2026, 1, 15), "alloy", 1, 10, 10, 10),
        (datetime(2026, 1, 16), "alloy", 1, 30, 30, 30),
    ]


@pytest.mark.parametrize("bounds", [
    {"end": "9999-12-31T23:59:59.999999Z"},
    {"end": "0001-01-01T00:00:00Z"},
    {"start": "9999-12-31T23:59:59.999999Z"},
])
async def test_realtime_extreme_range_is_bounded(client, accounts, bounds):
    response = await client.get("/api/v1/analytics/trend", headers=accounts.viewer.headers,
                                params={"source": "realtime", **bounds})
    assert response.status_code == 200, response.text
    assert response.json() == []


async def test_fall_back_cap_uses_utc_order(client, accounts, monkeypatch):
    from app.services import analytics

    # Local 01:05 GMT is later than local 01:55 BST, despite lexical local order.
    for instant in ["2026-10-25T00:55:00Z", "2026-10-25T01:05:00Z"]:
        await create_record(client, accounts.engineer, recorded_at=instant)
    monkeypatch.setattr(analytics, "MAX_TREND_POINTS", 1)
    response = await client.get("/api/v1/analytics/trend", headers=accounts.viewer.headers,
                                params={"tz": "Europe/London", "interval": "minute"})
    assert [point["bucket"] for point in response.json()] == ["2026-10-25T01:05:00Z"]


async def test_full_supported_record_span_still_caps_results(client, accounts, monkeypatch):
    from app.services import analytics

    for instant in ["1970-01-01T00:30:00Z", "2999-12-31T23:30:00Z"]:
        await create_record(client, accounts.engineer, recorded_at=instant)
    monkeypatch.setattr(analytics, "MAX_TREND_POINTS", 1)
    response = await client.get("/api/v1/analytics/trend", headers=accounts.viewer.headers,
                                params={"tz": "Europe/London", "interval": "hour"})
    assert response.status_code == 200, response.text
    assert [(point["bucket"], point["count"]) for point in response.json()] == [("2999-12-31T23:00:00Z", 1)]


@pytest.mark.parametrize("source", ["records", "realtime"])
@pytest.mark.parametrize("zone", ["Asia/Taipei", "America/Bogota"])
@pytest.mark.parametrize("boundary,instant", [
    ("start", "0001-01-01T00:00:00Z"), ("end", "0001-01-01T00:00:00Z"),
    ("start", "9999-12-31T23:59:59.999999Z"), ("end", "9999-12-31T23:59:59.999999Z"),
])
async def test_excel_extreme_range_labels_are_readable(client, accounts, source, zone, boundary, instant):
    response = await client.get("/api/v1/analytics/export", headers=accounts.viewer.headers,
                                params={"source": source, "tz": zone, boundary: instant})
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["count"] == 0
    assert len(list(workbook["Trend"].iter_rows(values_only=True))) == 1
    assert len(list(workbook["Data"].iter_rows(values_only=True))) == 1
    empty_minutes = source == "realtime" and ((boundary == "end" and instant.startswith("0001"))
                                               or (boundary == "start" and instant.startswith("9999")))
    if empty_minutes:
        assert summary["Range"].startswith("empty: no whole minutes")
    # Realtime end=max is capped to now; other labels retain the actual query bound.
    if not (source == "realtime" and boundary == "end" and instant.startswith("9999")):
        label = summary[boundary.capitalize()]
        assert label.startswith(instant[:4])
        if (zone == "America/Bogota" and instant.startswith("0001")) or (
            zone == "Asia/Taipei" and instant.startswith("9999")
        ):
            assert "UTC; outside" in label and zone in label
