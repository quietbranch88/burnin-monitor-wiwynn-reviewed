import io
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from app.core.time import utcnow
from app.db.session import SessionLocal
from app.models import RealtimeMinuteStat, RealtimeReading
from app.services import analytics
from openpyxl import load_workbook
from sqlalchemy import select

from tests.conftest import create_record

ROWS = [
    ("cpu_temp", 60.0, "2026-10-01T01:10:00Z"),
    ("cpu_temp", 80.0, "2026-10-01T01:50:00Z"),
    ("cpu_temp", 100.0, "2026-10-02T03:00:00Z"),
    ("power_w", 500.0, "2026-10-01T02:00:00Z"),
    ("power_w", 700.0, "2026-10-05T02:00:00Z"),
]


async def seed(client, accounts):
    for category, value, at in ROWS:
        await create_record(client, accounts.engineer, category=category, value=value, recorded_at=at)


async def test_summary_and_time_range(client, accounts):
    await seed(client, accounts)
    headers = accounts.viewer.headers
    everything = (await client.get("/api/v1/analytics/summary", params={"category": "cpu_temp"}, headers=headers)).json()
    assert everything == {"count": 3, "total": 240.0, "average": 80.0, "minimum": 60.0, "maximum": 100.0}
    window = await client.get(
        "/api/v1/analytics/summary",
        params={"start": "2026-10-01T00:00:00Z", "end": "2026-10-01T23:59:59Z"},
        headers=headers,
    )
    assert window.json()["count"] == 3 and window.json()["maximum"] == 500.0
    empty = await client.get("/api/v1/analytics/summary", params={"category": "none"}, headers=headers)
    assert empty.json() == {"count": 0, "total": None, "average": None, "minimum": None, "maximum": None}


async def test_aggregation_by_category(client, accounts):
    await seed(client, accounts)
    rows = (await client.get("/api/v1/analytics/categories", headers=accounts.viewer.headers)).json()
    assert [(r["category"], r["count"], r["average"]) for r in rows] == [("cpu_temp", 3, 80.0), ("power_w", 2, 600.0)]


async def test_trend_buckets(client, accounts):
    await seed(client, accounts)
    headers = accounts.viewer.headers
    hourly = (await client.get("/api/v1/analytics/trend", params={"interval": "hour", "category": "cpu_temp"}, headers=headers)).json()
    assert [(p["bucket"], p["count"], p["average"]) for p in hourly] == [
        ("2026-10-01T01:00:00Z", 2, 70.0),
        ("2026-10-02T03:00:00Z", 1, 100.0),
    ]
    daily = (await client.get("/api/v1/analytics/trend", params={"interval": "day"}, headers=headers)).json()
    assert [(p["bucket"][:10], p["category"], p["count"]) for p in daily] == [
        ("2026-10-01", "cpu_temp", 2), ("2026-10-01", "power_w", 1), ("2026-10-02", "cpu_temp", 1), ("2026-10-05", "power_w", 1),
    ]


async def test_realtime_source_reads_minute_statistics(client, accounts):
    async with SessionLocal() as session:
        session.add_all([
            RealtimeMinuteStat(minute=datetime(2026, 10, 1, 0, 0), category="cpu_temp", count=2, total=175.0,
                               minimum=80.0, maximum=95.0, anomalies=1),
            RealtimeMinuteStat(minute=datetime(2026, 10, 1, 0, 1), category="cpu_temp", count=2, total=150.0,
                               minimum=70.0, maximum=80.0, anomalies=0),
        ])
        await session.commit()
    headers = accounts.viewer.headers
    summary = (await client.get("/api/v1/analytics/summary", params={"source": "realtime"}, headers=headers)).json()
    assert summary == {"count": 4, "total": 325.0, "average": 81.25, "minimum": 70.0, "maximum": 95.0}
    minutes = (await client.get("/api/v1/analytics/trend", params={"source": "realtime", "interval": "minute"}, headers=headers)).json()
    assert [(p["bucket"], p["count"], p["average"]) for p in minutes] == [
        ("2026-10-01T00:00:00Z", 2, 87.5), ("2026-10-01T00:01:00Z", 2, 75.0),
    ]
    hourly = (await client.get("/api/v1/analytics/trend", params={"source": "realtime", "interval": "hour"}, headers=headers)).json()
    assert [(p["count"], p["maximum"]) for p in hourly] == [(4, 95.0)]


async def test_excel_export(client, accounts):
    await seed(client, accounts)
    response = await client.get("/api/v1/analytics/export", params={"category": "cpu_temp"}, headers=accounts.viewer.headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment; filename=" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "By category", "Trend", "Data"]
    data = list(workbook["Data"].iter_rows(values_only=True))
    assert data[0][:4] == ("id", "title", "value", "category")
    assert len(data) == 1 + 3
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["count"] == 3 and summary["average"] == 80.0
    by_category = list(workbook["By category"].iter_rows(values_only=True))
    assert by_category == [("category", "count", "total", "average", "minimum", "maximum"), ("cpu_temp", 3, 240.0, 80.0, 60.0, 100.0)]
    assert [row[2] for row in data[1:]] == [100.0, 80.0, 60.0]  # newest first, the whole set


async def test_excel_data_sheet_keeps_the_newest_rows_and_says_so(client, accounts, monkeypatch):
    await seed(client, accounts)
    monkeypatch.setattr(analytics, "EXPORT_ROW_LIMIT", 2)
    response = await client.get("/api/v1/analytics/export", params={"category": "cpu_temp"}, headers=accounts.viewer.headers)
    workbook = load_workbook(io.BytesIO(response.content))
    data = list(workbook["Data"].iter_rows(values_only=True))
    assert [row[2] for row in data[1:3]] == [100.0, 80.0] and data[3][0] == "Truncated to the latest 2 rows" and len(data) == 4
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["count"] == 3 and summary["Data sheet"].startswith("Only the latest 2 rows")


async def test_excel_neutralises_formulas(client, accounts):
    await create_record(client, accounts.engineer, title='=HYPERLINK("http://evil.example","x")')
    response = await client.get("/api/v1/analytics/export", headers=accounts.viewer.headers)
    workbook = load_workbook(io.BytesIO(response.content))
    title = workbook["Data"]["B2"]
    assert title.data_type == "s" and title.value == '=HYPERLINK("http://evil.example","x")'  # text, unchanged


async def test_trend_cap_keeps_newest_buckets(client, accounts, monkeypatch):
    await seed(client, accounts)
    monkeypatch.setattr(analytics, "MAX_TREND_POINTS", 2)
    points = (await client.get("/api/v1/analytics/trend", params={"interval": "day"}, headers=accounts.viewer.headers)).json()
    assert [p["bucket"][:10] for p in points] == ["2026-10-02", "2026-10-05"]


async def test_daily_trend_follows_the_requested_time_zone(client, accounts):
    # Both are on 21 September in UTC; in Taipei (UTC+8) the second is already 22 September.
    for at in ("2026-09-21T15:00:00Z", "2026-09-21T17:29:00Z"):
        await create_record(client, accounts.engineer, recorded_at=at)
    headers = accounts.viewer.headers
    utc = (await client.get("/api/v1/analytics/trend", params={"interval": "day"}, headers=headers)).json()
    assert [(p["bucket"], p["count"]) for p in utc] == [("2026-09-21T00:00:00Z", 2)]
    taipei = (await client.get("/api/v1/analytics/trend", params={"interval": "day", "tz": "Asia/Taipei"}, headers=headers)).json()
    # Buckets start at Taipei midnight, returned in UTC.
    assert [(p["bucket"], p["count"]) for p in taipei] == [("2026-09-20T16:00:00Z", 1), ("2026-09-21T16:00:00Z", 1)]
    bad = await client.get("/api/v1/analytics/trend", params={"tz": "Mars/Base"}, headers=headers)
    assert bad.status_code == 422


async def test_trend_shifts_negative_and_large_offsets(client, accounts):
    await create_record(client, accounts.engineer, recorded_at="2026-09-21T03:00:00Z")
    await create_record(client, accounts.engineer, recorded_at="1970-01-01T02:00:00Z")
    headers = accounts.viewer.headers

    async def days(zone: str) -> list[str]:
        params = {"interval": "day", "tz": zone}
        return [p["bucket"] for p in (await client.get("/api/v1/analytics/trend", params=params, headers=headers)).json()]

    # Bogota is UTC-5: 03:00 UTC is still the previous evening there.
    assert await days("America/Bogota") == ["1969-12-31T05:00:00Z", "2026-09-20T05:00:00Z"]
    # Kiritimati was UTC-10:40 in 1970; today it is UTC+14.
    assert await days("Pacific/Kiritimati") == ["1969-12-31T10:40:00Z", "2026-09-20T10:00:00Z"]


async def test_record_times_outside_the_supported_years_are_refused(client, accounts):
    for at in ("1969-12-31T23:59:59Z", "3000-01-01T00:00:00Z"):
        body = {"title": "t", "value": 1, "category": "cpu_temp", "recorded_at": at, "server": "S0001"}
        response = await client.post("/api/v1/records", json=body, headers=accounts.engineer.headers)
        assert response.status_code == 422 and "between 1970 and 2999" in response.text


async def test_excel_trend_uses_the_requested_time_zone(client, accounts):
    await create_record(client, accounts.engineer, recorded_at="2026-09-21T17:29:00Z")  # 22 September in Taipei
    params = {"interval": "day", "tz": "Asia/Taipei"}
    response = await client.get("/api/v1/analytics/export", params=params, headers=accounts.viewer.headers)
    trend = list(load_workbook(io.BytesIO(response.content))["Trend"].iter_rows(values_only=True))
    assert trend[0][0] == "bucket (Asia/Taipei)" and trend[1][0] == datetime(2026, 9, 22)


async def test_excel_days_start_at_midnight_in_the_other_daylight_saving_season(client, accounts):
    zone = ZoneInfo("Europe/London")
    now_offset = datetime.now(zone).utcoffset()
    # A date whose own offset differs from today's must still use historical midnight.
    day = next(d for d in (datetime(2026, 1, 15, 12), datetime(2026, 7, 15, 12))
               if d.replace(tzinfo=zone).utcoffset() != now_offset)
    await create_record(client, accounts.engineer, recorded_at=day.isoformat() + "Z")
    params = {"interval": "day", "tz": "Europe/London"}
    response = await client.get("/api/v1/analytics/export", params=params, headers=accounts.viewer.headers)
    trend = list(load_workbook(io.BytesIO(response.content))["Trend"].iter_rows(values_only=True))
    assert trend[1][0] == day.replace(hour=0)


async def test_one_server_statistics_come_from_its_raw_readings(client, accounts):

    async with SessionLocal() as session:
        # Fleet-wide minute statistics say nothing about one server...
        session.add(RealtimeMinuteStat(minute=datetime(2026, 10, 1, 0, 0), category="gpu_temp", count=3, total=210,
                                       minimum=60, maximum=80, anomalies=0))
        for server, value in (("S0001", 60.0), ("S0001", 70.0), ("S0002", 80.0)):
            session.add(RealtimeReading(server=server, category="gpu_temp", sensor="GPU0", value=value, threshold=85,
                                        is_anomaly=False, generated_at=datetime(2026, 10, 1, 0, 0, int(value) % 60)))
        await session.commit()
    params = {"source": "realtime", "start": "2026-10-01T00:00:00Z", "end": "2026-10-01T00:01:00Z"}
    headers = accounts.viewer.headers
    fleet = (await client.get("/api/v1/analytics/summary", params=params, headers=headers)).json()
    one = (await client.get("/api/v1/analytics/summary", params=params | {"server": "S0001"}, headers=headers)).json()
    assert fleet["count"] == 3 and (one["count"], one["average"], one["maximum"]) == (2, 65.0, 70.0)
    trend = (await client.get("/api/v1/analytics/trend", params=params | {"server": "S0001", "interval": "minute"},
                              headers=headers)).json()
    assert [(p["category"], p["count"]) for p in trend] == [("gpu_temp", 2)]
    report = await client.get("/api/v1/analytics/export", params=params | {"server": "S0001"}, headers=headers)
    workbook = load_workbook(io.BytesIO(report.content))
    servers = {row[1] for row in list(workbook["Data"].iter_rows(values_only=True))[1:] if row[0] is not None and isinstance(row[0], int)}
    assert servers == {"S0001"}
    records = (await client.get("/api/v1/analytics/summary", params={"server": "S0001"}, headers=headers)).json()
    assert records["count"] == 0  # no test records for that server in this test


async def test_one_server_range_is_bounded_because_it_scans_raw_readings(client, accounts):
    now = utcnow().replace(second=0, microsecond=0)
    async with SessionLocal() as session:
        for minutes_ago, value in ((30, 60.0), (120, 90.0)):
            session.add(RealtimeReading(server="S0001", category="gpu_temp", sensor="GPU0", value=value, threshold=85,
                                        is_anomaly=False, generated_at=now - timedelta(minutes=minutes_ago)))
        await session.commit()
    headers, params = accounts.viewer.headers, {"source": "realtime", "server": "S0001"}
    # Without a start: the last hour, so the reading of two hours ago is outside.
    assert (await client.get("/api/v1/analytics/summary", params=params, headers=headers)).json()["count"] == 1
    four_hours = {"start": (now - timedelta(hours=4)).isoformat() + "Z"}
    too_long = await client.get("/api/v1/analytics/summary", params=params | four_hours, headers=headers)
    assert too_long.status_code == 422 and "3 hours" in too_long.json()["detail"]
    # The fleet figures come from the per-minute statistics: no bound needed there.
    fleet = await client.get("/api/v1/analytics/summary", params={"source": "realtime", **four_hours}, headers=headers)
    assert fleet.status_code == 200


async def test_live_telemetry_covers_whole_minutes_for_the_fleet_and_one_server_alike(client, accounts):
    base = datetime(2026, 10, 1, 0, 0)
    async with SessionLocal() as session:
        for m in range(3):
            session.add(RealtimeMinuteStat(minute=base + timedelta(minutes=m), category="cpu_temp", count=2, total=160.0,
                                           minimum=70.0, maximum=90.0, anomalies=0))
            for second in (10, 50):
                session.add(RealtimeReading(server="S0001", category="cpu_temp", sensor="CPU0", value=80.0, threshold=90,
                                            is_anomaly=False, generated_at=base + timedelta(minutes=m, seconds=second)))
        await session.commit()

    async def count(**params):
        response = await client.get("/api/v1/analytics/summary", params={"source": "realtime", **params}, headers=accounts.viewer.headers)
        return response.json()["count"] if response.status_code == 200 else response.status_code

    # 00:00:30 to 00:02:30 holds one whole minute, 00:01: the partial minutes at both ends are not counted
    assert await count(start="2026-10-01T00:00:30Z", end="2026-10-01T00:02:30Z") == 2
    assert await count(start="2026-10-01T00:00:30Z", end="2026-10-01T00:02:30Z", server="S0001") == 2
    # An end on the minute closes the minute before it: 00:00 and 00:01, as a SQL BETWEEN would count
    assert await count(start="2026-10-01T00:00:00Z", end="2026-10-01T00:02:00Z") == 4
    assert await count(start="2026-10-01T00:00:00Z", end="2026-10-01T00:02:00Z", server="S0001") == 4
    assert await count(start="2026-10-01T00:00:00Z", end="2026-10-01T00:02:59.999999Z") == 6
    assert await count(start="2026-10-01T00:00:10Z", end="2026-10-01T00:00:50Z") == 0  # no whole minute inside
    assert await count(start="2026-10-01T00:02:00Z", end="2026-10-01T00:01:00Z") == 422  # reversed


async def test_minute_statistics_are_rebuilt_from_readings_after_an_unclean_stop():
    from app.services.realtime import backfill_minute_stats

    base = datetime(2026, 10, 1, 0, 0)
    started = base + timedelta(minutes=2, seconds=30)
    readings = [  # 00:00 had a partial row from a clean stop; 00:01 was lost in a kill; 00:02 spans this start
        (10, 70.0, False), (50, 95.0, True), (65, 80.0, False), (66, 85.0, False), (125, 60.0, False), (160, 99.0, True),
    ]
    async with SessionLocal() as session:
        session.add(RealtimeMinuteStat(minute=base, category="cpu_temp", count=1, total=70.0, minimum=70.0, maximum=70.0, anomalies=0))
        session.add_all(RealtimeReading(server="S0001", category="cpu_temp", sensor="CPU0", value=value, threshold=90,
                                        is_anomaly=anomaly, generated_at=base + timedelta(seconds=second))
                        for second, value, anomaly in readings)
        await session.commit()
    assert await backfill_minute_stats(started) == 2
    async with SessionLocal() as session:
        rows = {r.minute: r for r in await session.scalars(select(RealtimeMinuteStat))}
    assert (rows[base].count, rows[base].maximum, rows[base].anomalies) == (2, 95.0, 1)  # replaced, not added to
    assert (rows[base + timedelta(minutes=1)].count, rows[base + timedelta(minutes=1)].total) == (2, 165.0)  # the lost minute
    # Leave the startup minute alone: its generator may already have committed a contribution.
    assert base + timedelta(minutes=2) not in rows
    # A later recovery repairs it once the whole minute is outside the active window.
    assert await backfill_minute_stats(started + timedelta(minutes=1)) == 3
    async with SessionLocal() as session:
        recovered = await session.get(RealtimeMinuteStat, (base + timedelta(minutes=2), "cpu_temp"))
    assert recovered is not None and (recovered.count, recovered.maximum) == (2, 99.0)


async def test_excel_says_when_retention_has_removed_raw_readings(client, accounts):
    base = datetime(2026, 10, 1, 0, 0)
    async with SessionLocal() as session:
        session.add(RealtimeMinuteStat(minute=base, category="cpu_temp", count=4, total=320.0, minimum=70.0, maximum=90.0, anomalies=0))
        session.add(RealtimeReading(server="S0001", category="cpu_temp", sensor="CPU0", value=80.0, threshold=90,
                                    is_anomaly=False, generated_at=base + timedelta(seconds=10, microseconds=999999)))
        await session.commit()
    params = {"source": "realtime", "start": "2026-10-01T00:00:00Z", "end": "2026-10-01T00:00:59.999999Z"}
    response = await client.get("/api/v1/analytics/export", params=params | {"category": "cpu_temp"}, headers=accounts.viewer.headers)
    summary = {row[0]: row[1] for row in load_workbook(io.BytesIO(response.content))["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert summary["count"] == 4 and summary["Data sheet"].startswith("1 rows: retention has deleted")
    assert summary["Range"].startswith("whole minutes")
    workbook = load_workbook(io.BytesIO((await client.get("/api/v1/analytics/export", params=params, headers=accounts.viewer.headers)).content))
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    assert "units differ" in summary["total, average, minimum, maximum"]  # no category: a mix of units
    data = list(workbook["Data"].iter_rows(values_only=True))
    assert data[1][7] == datetime(2026, 10, 1, 0, 0, 10, 999000)  # cut to the millisecond, not rounded up
