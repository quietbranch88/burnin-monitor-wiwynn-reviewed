"""API contracts: validation, error bodies, request IDs, token revocation and report boundaries."""

import io
from datetime import timedelta

import pytest
from app.core.time import utcnow
from app.db.session import SessionLocal
from app.models import RealtimeMinuteStat, RealtimeReading
from openpyxl import load_workbook

from tests.conftest import create_record

RECORD = {"title": "t", "value": 1, "category": "cpu_temp", "recorded_at": "2026-10-01T08:00:00Z", "server": "S0001"}


@pytest.mark.parametrize("value", [1.7e308, -2e15, True, False])
async def test_values_that_cannot_be_summed_or_are_not_numbers_are_refused(client, accounts, value):
    response = await client.post("/api/v1/records", json=RECORD | {"value": value}, headers=accounts.engineer.headers)
    assert response.status_code == 422


async def test_an_explicit_null_in_an_update_is_refused(client, accounts):
    record = await create_record(client, accounts.engineer)
    response = await client.patch(f"/api/v1/records/{record['id']}", json={"title": None}, headers=accounts.engineer.headers)
    assert response.status_code == 422 and "must not be null" in response.text


@pytest.mark.parametrize(("password", "accepted"), [
    ("aaaaaaaa", False), ("12345678", False), ("password", False), ("password1", False), ("aaaa1111", False),
    ("S3cure-pass", True),
])
async def test_trivial_passwords_are_refused(client, password, accepted):
    body = {"username": "new.user", "email": "new.user@example.com", "password": password}
    response = await client.post("/api/v1/auth/register", json=body)
    assert (response.status_code == 201) is accepted


async def test_openapi_documents_the_real_error_body(app):
    schema = app.openapi()["components"]["schemas"]["HTTPValidationError"]
    assert set(schema["properties"]) == {"detail", "request_id", "errors"}
    assert "ValidationError" not in app.openapi()["components"]["schemas"]
    assert "ws/realtime" in app.openapi()["info"]["description"]


async def test_a_body_that_is_not_json_names_the_body(client, accounts):
    headers = accounts.engineer.headers | {"Content-Type": "application/json"}
    response = await client.post("/api/v1/records", content=b"{not json", headers=headers)
    assert response.status_code == 422 and response.json()["errors"][0]["field"] == "body"


async def test_a_created_record_has_a_location(client, accounts):
    response = await client.post("/api/v1/records", json=RECORD, headers=accounts.engineer.headers)
    assert response.headers["location"] == f"/api/v1/records/{response.json()['id']}"


@pytest.mark.parametrize(("given", "kept"), [("trace-42.a:b", True), ("<script>", False), ("x" * 37, False)])
async def test_only_a_plain_request_id_is_kept(client, given, kept):
    response = await client.get("/health", headers={"X-Request-ID": given})
    assert (response.headers["x-request-id"] == given) is kept


async def test_event_logs_carry_the_request_id_and_can_be_filtered(app, client, accounts):
    path = f"/api/v1/admin/users/{accounts.viewer.user.id}"
    headers = accounts.admin.headers | {"X-Request-ID": "test-role-change-1"}
    assert (await client.patch(path, json={"role": "user"}, headers=headers)).status_code == 200
    await app.state.log_writer.flush()
    logs = (await client.get("/api/v1/admin/logs", params={"request_id": "test-role-change-1"},
                             headers=accounts.admin.headers)).json()["items"]
    assert {e["logger"] for e in logs} == {"app.users", "app.request"}
    users_only = (await client.get("/api/v1/admin/logs", params={"logger": "app.users"}, headers=accounts.admin.headers)).json()
    assert users_only["total"] >= 1 and {e["logger"] for e in users_only["items"]} == {"app.users"}


async def test_deactivating_ends_the_tokens_even_after_reactivation(client, accounts):
    path = f"/api/v1/admin/users/{accounts.viewer.user.id}"
    for active in (False, True):
        assert (await client.patch(path, json={"is_active": active}, headers=accounts.admin.headers)).status_code == 200
    assert (await client.get("/api/v1/auth/me", headers=accounts.viewer.headers)).status_code == 401
    refused = await client.patch(path, json={"is_active": "no"}, headers=accounts.admin.headers)
    assert refused.status_code == 422


async def test_a_csv_error_names_the_line_it_is_on(client, accounts):
    content = b"title,value,category,recorded_at,server\nok,1,cpu_temp,2026-10-01T08:00:00Z,S0001\n" + b"x" * 200_000 + b",1,cpu_temp,x,S0001\n"
    files = {"file": ("rows.csv", content, "text/csv")}
    response = await client.post("/api/v1/records/import", files=files, headers=accounts.engineer.headers)
    assert response.status_code == 422 and "line 3" in response.json()["detail"]


async def test_live_telemetry_data_sheet_covers_the_same_minutes_as_the_summary(client, accounts):
    minute = utcnow().replace(second=0, microsecond=0) - timedelta(minutes=10)
    async with SessionLocal() as session:
        for m in range(3):
            start = minute + timedelta(minutes=m)
            session.add(RealtimeMinuteStat(minute=start, category="cpu_temp", count=2, total=160, minimum=70, maximum=90, anomalies=0))
            for second in (10, 40):
                session.add(RealtimeReading(server=f"S{m}{second}", category="cpu_temp", sensor="CPU0", value=80, threshold=90,
                                            is_anomaly=False, generated_at=start + timedelta(seconds=second)))
        await session.commit()
    # 20 s into the first minute until 5 s into the last one: only the middle minute lies whole
    # inside the range, so the summary counts it alone, and the Data sheet lists the same readings.
    params = {"source": "realtime", "category": "cpu_temp",
              "start": (minute + timedelta(seconds=20)).isoformat() + "Z",
              "end": (minute + timedelta(minutes=2, seconds=5)).isoformat() + "Z"}
    response = await client.get("/api/v1/analytics/export", params=params, headers=accounts.viewer.headers)
    workbook = load_workbook(io.BytesIO(response.content))
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    header, *data_rows = list(workbook["Data"].iter_rows(values_only=True))
    assert summary["count"] == 2 and len(data_rows) == 2
    assert header[:4] == ("id", "server", "category", "sensor") and {row[3] for row in data_rows} == {"CPU0"}


async def test_role_polling_does_not_flood_the_log_and_entries_name_the_user(app, client, accounts):
    for n in range(3):
        headers = accounts.viewer.headers | {"X-Request-ID": f"test-poll-ok-{n}"}
        assert (await client.get("/api/v1/auth/me", headers=headers)).status_code == 200
    assert (await client.get("/api/v1/auth/me", headers={"X-Request-ID": "test-poll-fail"})).status_code == 401
    await client.get("/api/v1/records", headers=accounts.engineer.headers | {"X-Request-ID": "test-records-1"})
    await app.state.log_writer.flush()

    async def lines(request_id: str) -> list[dict]:
        params = {"request_id": request_id}
        return (await client.get("/api/v1/admin/logs", params=params, headers=accounts.admin.headers)).json()["items"]

    for n in range(3):
        assert await lines(f"test-poll-ok-{n}") == []  # successful polls are not logged
    assert [e["status_code"] for e in await lines("test-poll-fail")] == [401]
    assert (await lines("test-records-1"))[0]["username"] == "engineer"


@pytest.mark.parametrize("threshold", [45.0, 75.0, 3.0])
def test_simulated_alert_rate_follows_any_configured_threshold(threshold):
    import random

    from app.services.realtime import generate_value

    rng = random.Random(7)
    alerts = sum(generate_value("cpu_temp", threshold, rng) > threshold for _ in range(200_000))
    assert 0.0001 < alerts / 200_000 < 0.0006


async def test_a_seed_password_that_breaks_the_rule_stops_with_a_clear_message(monkeypatch):
    from app import seed
    from app.core.config import get_settings

    monkeypatch.setattr(seed, "DEMO_ACCOUNTS", [("seedcheck", "seedcheck@example.com", seed.Role.viewer, "seed_viewer_password")])
    monkeypatch.setattr(get_settings(), "seed_viewer_password", "aaaaaaaa")
    monkeypatch.setattr(seed, "engine", type("E", (), {"dispose": staticmethod(lambda: None)}))
    with pytest.raises(SystemExit, match="SEED_VIEWER_PASSWORD is not accepted"):
        await seed.seed()


async def test_a_missing_body_names_the_body(client, accounts):
    response = await client.post("/api/v1/records", headers=accounts.engineer.headers)
    assert response.status_code == 422 and response.json()["errors"][0]["field"] == "body"


async def test_admins_can_search_users(client, accounts):
    found = (await client.get("/api/v1/admin/users", params={"q": "engin"}, headers=accounts.admin.headers)).json()
    assert found["total"] == 1 and found["items"][0]["username"] == "engineer"
    by_email = (await client.get("/api/v1/admin/users", params={"q": "viewer@"}, headers=accounts.admin.headers)).json()
    assert [u["username"] for u in by_email["items"]] == ["viewer"]


async def test_live_telemetry_data_sheet_leaves_out_the_minute_in_progress(client, accounts):
    now = utcnow()
    finished = now.replace(second=0, microsecond=0) - timedelta(minutes=1)
    async with SessionLocal() as session:
        session.add(RealtimeMinuteStat(minute=finished, category="cpu_temp", count=1, total=80, minimum=80, maximum=80, anomalies=0))
        for at in (finished + timedelta(seconds=30), now):  # the second one has no statistics yet
            session.add(RealtimeReading(server="S0001", category="cpu_temp", sensor="CPU0", value=80, threshold=90,
                                        is_anomaly=False, generated_at=at))
        await session.commit()
    params = {"source": "realtime", "start": (finished - timedelta(minutes=5)).isoformat() + "Z"}  # no end: up to now
    response = await client.get("/api/v1/analytics/export", params=params, headers=accounts.viewer.headers)
    workbook = load_workbook(io.BytesIO(response.content))
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row and row[0]}
    data_rows = list(workbook["Data"].iter_rows(values_only=True))[1:]
    assert summary["count"] == 1 and len(data_rows) == 1
